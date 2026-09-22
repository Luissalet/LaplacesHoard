"""Regressions from the first user walk (docs/USABILITY_REPORT.md): locale
CSV handling, re-registration, and nested/BLOB values blowing up a result."""
import csv
import json
from pathlib import Path

import openpyxl
import pytest

from laplaces_hoard.engines.data import Catalog, DataError


def _write(path: Path, lines: list[str], encoding: str = "utf-8") -> Path:
    with path.open("w", encoding=encoding, newline="") as fh:
        for line in lines:
            fh.write(line + "\n")
    return path


def test_spanish_decimal_and_thousands_separators_are_cast_to_numbers(tmp_path):
    p = _write(tmp_path / "ventas.csv", ["fecha,importe", '13/02/25,"-1.150,00"', '01/12/24,"2.300,50"'])
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p), options={"decimal_separator": ",", "thousands_separator": "."})
    types = {c["name"]: c["type"] for c in meta["columns"]}
    assert types["importe"] == "DOUBLE"
    row = cat.query("SELECT SUM(importe) AS s FROM ventas")["rows"][0]
    assert row["s"] == pytest.approx(1150.5)


def test_decimal_separator_alone_does_not_touch_genuinely_textual_columns(tmp_path):
    # a comma-decimal column next to a plain text column: only the numeric
    # one should be rewritten, the id/text columns stay as they were
    p = _write(tmp_path / "t.csv", ["id,label,amount", '1,North,"1,50"', '2,South,"2,75"'])
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p), options={"decimal_separator": ","})
    types = {c["name"]: c["type"] for c in meta["columns"]}
    assert types["amount"] == "DOUBLE"
    assert types["label"] == "VARCHAR"
    assert cat.query("SELECT label FROM t WHERE id = 1")["rows"][0]["label"] == "North"


def test_two_digit_year_dates_are_read_day_first_when_unambiguous(tmp_path):
    # "13/02/25": 13 can only be a day, so this must be 2025-02-13, not
    # 2013-02-25 (DuckDB's own sniffer's default reading of this exact input)
    p = _write(tmp_path / "t.csv", ["fecha", "13/02/25", "01/12/24"])
    cat = Catalog(tmp_path / "data")
    cat.register(str(p))
    rows = cat.query("SELECT fecha FROM t ORDER BY fecha")["rows"]
    assert [r["fecha"] for r in rows] == ["2024-12-01", "2025-02-13"]


def test_explicit_date_format_overrides_the_sniffer(tmp_path):
    p = _write(tmp_path / "t.csv", ["fecha", "01/02/2024"])
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p), options={"date_format": "%m/%d/%Y"})
    assert cat.query("SELECT fecha FROM t")["rows"][0]["fecha"] == "2024-01-02"
    assert meta["row_count"] == 1


def test_ambiguous_dates_are_left_to_duckdbs_own_sniffer(tmp_path):
    # every component <= 12: no evidence either way, so behaviour is
    # unchanged from before this fix (not a promise it's "right", just that
    # we do not add a new wrong guess on top of an old one)
    p = _write(tmp_path / "t.csv", ["fecha", "01/02/2024"])
    cat = Catalog(tmp_path / "data")
    cat.register(str(p))
    # DuckDB's default sniffer reads this as month/day here; the point of
    # this test is only that registration succeeds and is deterministic.
    row = cat.query("SELECT fecha FROM t")["rows"][0]
    assert row["fecha"] in ("2024-02-01", "2024-01-02")


def test_windows_1252_csv_is_registered_without_an_explicit_option(tmp_path):
    p = tmp_path / "clientes.csv"
    with p.open("wb") as fh:
        fh.write("nombre,ciudad\n".encode("utf-8"))
        fh.write("José,León\n".encode("cp1252"))
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p))
    assert meta["encoding_detected"] == "latin-1"
    assert meta["sample_rows"][0] == {"nombre": "José", "ciudad": "León"}


def test_explicit_encoding_option_is_honoured(tmp_path):
    p = tmp_path / "t.csv"
    with p.open("wb") as fh:
        fh.write("a\n".encode("utf-8"))
        fh.write("café\n".encode("latin-1"))
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p), options={"encoding": "latin-1"})
    assert meta["sample_rows"][0]["a"] == "café"


def test_unsupported_encoding_option_is_a_clear_error(tmp_path):
    p = _write(tmp_path / "t.csv", ["a", "1"])
    cat = Catalog(tmp_path / "data")
    with pytest.raises(DataError, match="unsupported encoding"):
        cat.register(str(p), options={"encoding": "shift-jis"})


def test_registering_the_same_csv_twice_refreshes_it(tmp_path):
    # the documented way to pick up a changed file; used to fail with
    # "existing object ... is of type Table, trying to drop type View"
    p = _write(tmp_path / "t.csv", ["a,b", "1,2"])
    cat = Catalog(tmp_path / "data")
    cat.register(str(p))
    _write(p, ["a,b", "1,2", "3,4"])
    meta = cat.register(str(p))
    assert meta["row_count"] == 2
    meta_again = cat.register(str(p))
    assert meta_again["row_count"] == 2


def test_registering_the_same_linked_big_file_twice_refreshes_it(tmp_path, monkeypatch):
    # the opposite direction: a VIEW being re-created where a VIEW already exists
    # is the easy case, but must still work after the generic drop-both fix
    from laplaces_hoard.engines import data as data_mod

    monkeypatch.setattr(data_mod, "LINK_THRESHOLD_BYTES", 0)
    p = _write(tmp_path / "t.csv", ["a", "1", "2"])
    cat = Catalog(tmp_path / "data")
    cat.register(str(p))
    meta = cat.register(str(p))
    assert meta["linked"] is True
    assert meta["row_count"] == 2


def test_single_row_with_a_nested_list_is_capped_and_hinted(tmp_path):
    p = tmp_path / "funes.json"
    row = {"title": "log", "events": [{"t": i, "app": f"App{i}"} for i in range(500)]}
    p.write_text(json.dumps([row]), encoding="utf-8")
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p))
    assert len(meta["sample_rows"][0]["events"]) <= 21  # 20 items + a "N more" marker
    assert "events" in meta["hint"]
    assert "UNNEST" in meta["hint"]
    assert len(json.dumps(meta)) < 20_000


def test_nested_list_in_a_public_query_result_is_also_capped(tmp_path):
    p = tmp_path / "funes.json"
    row = {"events": list(range(200))}
    p.write_text(json.dumps([row]), encoding="utf-8")
    cat = Catalog(tmp_path / "data")
    cat.register(str(p), name="funes")
    r = cat.query("SELECT events FROM funes")
    assert len(r["rows"][0]["events"]) <= 21
    # query_all (used by stats/charts internally) must NOT be capped: it
    # needs the real values, not a preview
    full = cat.query_all("SELECT events FROM funes")
    assert len(full["rows"][0]["events"]) == 200


def test_excel_title_row_above_the_header_is_detected_and_skipped(tmp_path):
    # a lone title in row 1 ("Informe trimestral Q1 2024") used to become
    # the header, producing columns named col0/col1 instead of the real ones
    xlsx = tmp_path / "informe.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Informe trimestral Q1 2024"])
    ws.append(["region", "total"])
    ws.append(["North", 100])
    ws.append(["South", 200])
    wb.save(xlsx)
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(xlsx))
    ds = meta["datasets"][0]
    assert ds["columns"] == ["region", "total"]
    assert ds["row_count"] == 2


def test_excel_without_a_title_row_is_unaffected(tmp_path):
    xlsx = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["region", "total"])
    ws.append(["North", 100])
    wb.save(xlsx)
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(xlsx))
    assert meta["datasets"][0]["columns"] == ["region", "total"]


def test_excel_skip_rows_option_overrides_auto_detection(tmp_path):
    xlsx = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["note 1"])
    ws.append(["note 2"])
    ws.append(["region", "total"])
    ws.append(["North", 100])
    wb.save(xlsx)
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(xlsx), options={"skip_rows": 2})
    assert meta["datasets"][0]["columns"] == ["region", "total"]


def test_sqlite_blob_is_a_real_blob_shown_as_its_size(tmp_path):
    # a thumbnail BLOB used to arrive as its Python repr text (b"\\x..."),
    # typed VARCHAR, flooding describe's top values and sample rows
    import sqlite3

    db_path = tmp_path / "b.sqlite"
    conn = sqlite3.connect(db_path)
    conn.execute("CREATE TABLE t (id INTEGER, data BLOB)")
    conn.execute("INSERT INTO t VALUES (1, ?)", (b"x" * 5000,))
    conn.execute("INSERT INTO t VALUES (2, ?)", (bytes(range(256)),))
    conn.commit()
    conn.close()
    cat = Catalog(tmp_path / "data")
    cat.register(str(db_path))
    d = cat.describe("b__t")
    profile = d["profile"]["data"]
    assert profile["type"] == "BLOB"
    assert "top_values" not in profile
    assert (profile["bytes_min"], profile["bytes_max"]) == (256, 5000)
    assert d["sample_rows"][0]["data"] == "<binary, 5,000 bytes>"
    assert len(json.dumps(d)) < 2000
    # the bytes themselves survive the import intact
    full = cat.query_all("SELECT data FROM b__t ORDER BY id")["rows"]
    assert full[1]["data"] == bytes(range(256)).hex()
    assert cat.query("SELECT octet_length(data) AS n FROM b__t ORDER BY id")["rows"][0]["n"] == 5000


# -- second walk (re-walk after the fix pass) ---------------------------------

def test_spanish_bank_amounts_become_exact_numbers_without_any_option(tmp_path):
    # the person pastes a path and presses Register: no option to set, and
    # SUM() must work straight away, exactly (money), not as 1150.4999...
    p = _write(tmp_path / "movimientos.csv", [
        "\ufeffFecha;Concepto;Importe (€);Saldo (€)",
        "01/07/2023;RECIBO ALQUILER;-1.150,00;7.262,37",
        "02/07/2023;COMPRA;-51,05;7.211,32",
        "03/07/2023;NOMINA;2.300,10;9.511,42",
        "04/07/2023;SIN IMPORTE;;9.511,42",
    ])
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p))
    types = {c["name"]: c["type"] for c in meta["columns"]}
    assert types["Importe (€)"] == "DECIMAL(18,2)"
    assert types["Saldo (€)"] == "DECIMAL(18,2)"
    assert types["Concepto"] == "VARCHAR"
    assert meta["numbers_converted"]["columns"] == ["Importe (€)", "Saldo (€)"]
    row = cat.query('SELECT SUM("Importe (€)") AS s, COUNT("Importe (€)") AS n FROM movimientos')["rows"][0]
    assert row["s"] == "1099.05"  # exact decimal, as text like every DECIMAL result
    assert row["n"] == 3  # the blank amount stays NULL


def test_automatic_detection_leaves_ambiguous_or_mixed_text_alone(tmp_path):
    p = _write(tmp_path / "mixto.csv", [
        "id;english_thousands;salario;version;codigo",
        '1;"1,234";55.000 €;1,2,3;0012',
        '2;"12,500";60-65k;2,0,1;0034',
    ])
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p))
    types = {c["name"]: c["type"] for c in meta["columns"]}
    # "1,234"/"12,500" could be English thousands: never guessed
    assert types["english_thousands"] == "VARCHAR"
    assert types["salario"] == "VARCHAR"
    assert types["version"] == "VARCHAR"
    assert "numbers_converted" not in meta


def test_decimal_separator_dot_opts_out_of_the_automatic_detection(tmp_path):
    p = _write(tmp_path / "keep.csv", ["a;b", "x;1,50", "y;2,25"])
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p), options={"decimal_separator": "."})
    assert {c["name"]: c["type"] for c in meta["columns"]}["b"] == "VARCHAR"


def test_excel_title_and_blank_spacer_row_above_the_header_are_skipped(tmp_path):
    # the job-hunt workbook's second sheet: a title, an empty row, then the header
    xlsx = tmp_path / "busqueda.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Entrevistas"
    ws.append(["Seguimiento de entrevistas 2026"])
    ws.append([])
    ws.append(["Empresa", "Fecha", "Fase"])
    ws.append(["Acme", "2026-05-01", "Técnica"])
    wb.save(xlsx)
    cat = Catalog(tmp_path / "data")
    ds = cat.register(str(xlsx))["datasets"][0]
    assert ds["columns"] == ["Empresa", "Fecha", "Fase"]
    assert ds["row_count"] == 1


def test_single_column_excel_sheet_is_never_skipped(tmp_path):
    xlsx = tmp_path / "lista.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for v in ("nombre", "Ana", "Luis"):
        ws.append([v])
    wb.save(xlsx)
    cat = Catalog(tmp_path / "data")
    ds = cat.register(str(xlsx))["datasets"][0]
    assert ds["columns"] == ["nombre"] and ds["row_count"] == 2


def test_nested_export_does_not_flood_register_or_describe(tmp_path):
    # a Funes-style export: one row, one big array of structs. Registering it
    # used to return ~200k characters (top_values of the whole array)
    p = tmp_path / "export.json"
    spans = [{"id": i, "start_ts": 1.0 * i, "end_ts": i + 0.5, "title": "Visual Studio Code - " + "x" * 80,
              "category": "Coding"} for i in range(400)]
    p.write_text(json.dumps({"spans": spans, "commits": [], "exported_at": 1.0}), encoding="utf-8")
    cat = Catalog(tmp_path / "data")
    meta = cat.register(str(p))
    assert len(json.dumps(meta)) < 4000
    assert "top_values" not in meta["profile"]["spans"]
    assert "UNNEST" in meta["hint"]
    r = cat.query("SELECT SUM(s.end_ts - s.start_ts) AS t FROM (SELECT UNNEST(spans) AS s FROM export)")
    assert r["rows"][0]["t"] == pytest.approx(200.0)
