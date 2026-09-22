#!/usr/bin/env python3
"""Generate the realistic (synthetic) inputs used by docs/USE_CASES.md.

Everything here is invented: no real person, account, company or photo.
The files imitate what a typical user in Spain has on disk - a
Spanish bank export (';' separators, decimal commas, dd/mm/yyyy dates,
Windows-1252, a preamble above the table), a job-hunt workbook, a writing
log, a llama.cpp benchmark log, a Funes activity export and an Daguerre
photo library - so the use cases are walked against messy input rather
than the tidy demo files.

Usage:
    .venv/bin/python scripts/uxtest_data.py [out_dir]   (default data-uxtest/files)
"""
from __future__ import annotations

import csv
import json
import random
import sqlite3
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
rng = random.Random(20260922)


def es_number(value: float, thousands: bool = True) -> str:
    """1234.5 -> '1.234,50' (Spanish style)."""
    text = f"{value:,.2f}" if thousands else f"{value:.2f}"
    return text.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


# ---------------------------------------------------------------- bank --

MERCHANTS = {
    "Supermercado": ["COMPRA MERCADONA C/ ALCALA", "COMPRA CARREFOUR EXPRESS", "COMPRA DIA %s", "COMPRA LIDL VALLECAS",
                     "COMPRA AHORRAMAS"],
    "Restaurantes": ["PAGO BAR LA CANA", "PAGO TABERNA EL SUR", "PAGO GLOVO*PEDIDO", "PAGO JUST EAT", "PAGO CAFETERIA NIZA"],
    "Transporte": ["RECARGA ABONO CRTM", "PAGO CABIFY", "PAGO RENFE CERCANIAS", "GASOLINERA REPSOL M-30"],
    "Suscripciones": ["CARGO SPOTIFY", "CARGO NETFLIX", "CARGO GITHUB", "CARGO ICLOUD"],
    "Hogar": ["RECIBO IBERDROLA", "RECIBO CANAL ISABEL II", "RECIBO DIGI FIBRA", "COMPRA IKEA SAN SEBASTIAN"],
    "Ocio": ["COMPRA FNAC CALLAO", "PAGO CINES GOLEM", "COMPRA STEAM", "PAGO TEATRO KAMIKAZE"],
    "Salud": ["FARMACIA LDO. PEREZ", "PAGO FISIOTERAPIA", "PAGO DENTISTA"],
    "Tecnología": ["COMPRA PCCOMPONENTES", "COMPRA AMAZON.ES", "COMPRA COOLMOD"],
}
AMOUNT_RANGE = {
    "Supermercado": (6, 95), "Restaurantes": (3, 45), "Transporte": (1.5, 55), "Suscripciones": (4.99, 15.99),
    "Hogar": (25, 180), "Ocio": (8, 90), "Salud": (4, 120), "Tecnología": (15, 1450),
}


def bank_rows() -> list[dict]:
    rows = []
    balance = 8_412.37
    day = date(2023, 7, 1)
    while day <= date(2025, 12, 31):
        if day.day == 28:  # payroll
            amt = round(rng.uniform(3890, 3960), 2)
            rows.append({"date": day, "concept": "TRANSFERENCIA NOMINA ACME SOFTWARE SL", "cat": "Nómina", "amount": amt})
        if day.day == 1:
            rows.append({"date": day, "concept": "RECIBO ALQUILER VIVIENDA", "cat": "Hogar", "amount": -1150.00})
        for _ in range(rng.choice([0, 1, 1, 2, 2, 3, 3, 4])):
            cat = rng.choices(list(MERCHANTS), weights=[30, 22, 18, 4, 6, 8, 4, 3])[0]
            lo, hi = AMOUNT_RANGE[cat]
            amt = -round(rng.uniform(lo, hi) if cat != "Tecnología" else rng.lognormvariate(3.6, 1.0), 2)
            amt = max(amt, -1899.0)
            concept = rng.choice(MERCHANTS[cat])
            if "%s" in concept:
                concept = concept % rng.choice(["LAVAPIES", "CHAMBERI", "ARGANZUELA"])
            rows.append({"date": day, "concept": concept, "cat": cat, "amount": amt})
        if rng.random() < 0.02:  # a refund
            rows.append({"date": day, "concept": "DEVOLUCION AMAZON.ES", "cat": "Tecnología", "amount": round(rng.uniform(15, 90), 2)})
        day += timedelta(days=1)
    for r in rows:
        balance = round(balance + r["amount"], 2)
        r["balance"] = balance
        r["value_date"] = r["date"] + timedelta(days=rng.choice([0, 0, 0, 1, 2]))
        if rng.random() < 0.06:
            r["cat"] = ""  # the bank did not classify it
    return rows


def write_bank(out: Path) -> None:
    rows = bank_rows()
    # 1) the "downloaded from online banking" version: UTF-8 with BOM, ';', decimal comma, dd/mm/yyyy
    with (out / "movimientos_2023-2025.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["Fecha operación", "Fecha valor", "Concepto", "Categoría", "Importe (€)", "Saldo (€)"])
        for r in rows:
            w.writerow([r["date"].strftime("%d/%m/%Y"), r["value_date"].strftime("%d/%m/%Y"), r["concept"], r["cat"],
                        es_number(r["amount"]), es_number(r["balance"])])
    # 2) the "export to Excel/CSV" of another bank: Windows-1252, preamble lines, no thousands dot
    with (out / "extracto_cuenta_2025.csv").open("w", encoding="cp1252", newline="") as fh:
        fh.write("Extracto de movimientos\r\n")
        fh.write("Cuenta;ES00 0000 0000 0000 0000 0000\r\n")
        fh.write(f"Periodo;{rows[-1500]['date']:%d/%m/%Y} - 31/12/2025\r\n")
        fh.write("\r\n")
        w = csv.writer(fh, delimiter=";", lineterminator="\r\n")
        w.writerow(["Fecha", "Concepto", "Importe", "Divisa", "Saldo"])
        for r in rows[-1500:]:
            w.writerow([r["date"].strftime("%d/%m/%Y"), r["concept"], es_number(r["amount"], thousands=False), "EUR",
                        es_number(r["balance"], thousands=False)])


# ---------------------------------------------------------------- jobs --

COMPANIES = ["Nimbus Labs", "Orbital Data", "Quark Systems", "Helios AI", "Vértice Software", "Cobalto Tech",
             "Lumen Robotics", "Faro Analytics", "Tesela Cloud", "Brújula Mobility", "Atlas Fintech", "Nébula Games",
             "Pixel Forge", "Río Health", "Sierra Security", "Delta Compilers", "Kora Media", "Monte Energy"]
ROLES = ["Senior Backend Engineer", "Staff Engineer", "ML Engineer", "Senior Python Developer", "Tech Lead",
         "Platform Engineer", "LLM Infrastructure Engineer", "Senior Software Engineer (C#/.NET)"]


def write_jobs(out: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aplicaciones"
    ws.append(["Empresa", "Puesto", "Fecha envío", "Modalidad", "Salario bruto", "Fuente", "Estado", "Notas"])
    start = date(2026, 5, 4)
    apps = []
    for i in range(96):
        d = start + timedelta(days=int(i * 1.3) + rng.choice([0, 0, 1]))
        mode = rng.choices(["Remoto", "Híbrido", "Presencial"], weights=[45, 40, 15])[0]
        p_reply = {"Remoto": 0.22, "Híbrido": 0.38, "Presencial": 0.45}[mode]
        replied = rng.random() < p_reply
        state = rng.choice(["Entrevista", "Rechazada", "Entrevista", "Oferta"]) if replied else rng.choice(
            ["Sin respuesta", "Sin respuesta", "Rechazada automática"])
        salary = rng.choice([None, 52000, 55000, 58000, 60000, 62000, 65000, 70000, "55.000 €", "60-65k", "A convenir"])
        company = rng.choice(COMPANIES)
        apps.append((company, d, state))
        ws.append([company, rng.choice(ROLES), d, mode, salary,
                   rng.choice(["LinkedIn", "InfoJobs", "Referido", "Web empresa", "Tecnoempleo"]), state,
                   rng.choice([None, None, "recruiter muy majo", "prueba técnica 4h", "pide inglés C1", None])])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = "DD/MM/YYYY"
    ws2 = wb.create_sheet("Entrevistas")
    ws2.append(["Seguimiento de entrevistas 2026"])  # a title row above the header, as people do
    ws2.append([])
    ws2.append(["Empresa", "Fecha", "Fase", "Duración (min)", "Sensación (1-5)"])
    for company, d, state in apps:
        if state in ("Entrevista", "Oferta"):
            for k, phase in enumerate(["Screening", "Técnica", "Cultural"][: rng.randint(1, 3)]):
                ws2.append([company, d + timedelta(days=7 + 6 * k), phase, rng.choice([30, 45, 60, 90]), rng.randint(2, 5)])
    wb.save(out / "busqueda_empleo_2026.xlsx")


# ------------------------------------------------------------- writing --

def write_writing(out: Path) -> None:
    with (out / "sesiones_escritura.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["date", "project", "chapter", "words", "minutes", "place"])
        d = date(2025, 10, 1)
        chapter = 1
        written_in_ch = 0
        while d <= date(2026, 9, 20):
            if rng.random() < 0.45:
                minutes = rng.choice([25, 30, 45, 50, 60, 75, 90, 120])
                words = int(minutes * rng.uniform(5, 12))
                place = rng.choice(["casa", "casa", "biblioteca", "tren", "café"])
                w.writerow([d.isoformat(), "Novela (borrador)", chapter, words, minutes, place])
                written_in_ch += words
                if written_in_ch > rng.randint(3200, 4600):
                    chapter += 1
                    written_in_ch = 0
            d += timedelta(days=1)


# ----------------------------------------------------------- llama.cpp --

def write_bench(out: Path) -> None:
    models = [("qwen3-27b-q4_k_m", 27, 17.5), ("qwen3-8b-q6_k", 8, 48.0), ("gemma3-12b-q4_k_m", 12, 33.0)]
    with (out / "llama_bench.ndjson").open("w", encoding="utf-8") as fh:
        t0 = datetime(2026, 9, 1, 21, 0)
        for i in range(1200):
            name, params, base = rng.choice(models)
            ctx = rng.choice([2048, 4096, 8192, 16384, 32768])
            tps = base * (1 - 0.0000085 * ctx) * rng.uniform(0.9, 1.08)
            rec = {"ts": (t0 + timedelta(minutes=7 * i)).isoformat(timespec="seconds"), "model": name,
                   "params_b": params, "ctx": ctx, "n_gen": rng.choice([128, 256, 512]),
                   "tokens_per_s": round(tps, 2), "gpu_temp_c": round(rng.uniform(58, 83), 1),
                   "vram_mb": int(params * 620 + ctx * 0.11 + rng.uniform(-150, 150))}
            if rng.random() < 0.01:
                rec["tokens_per_s"] = None  # a crashed run
            fh.write(json.dumps(rec) + "\n")


# ---------------------------------------------------------------- Funes --

APPS = [
    ("Code", "Code.exe", "Coding", ["api.py - laplaces-hoard - Visual Studio Code", "DataPage.tsx - laplaces-hoard - Visual Studio Code",
                                     "test_calc.py - babels-hoard - Visual Studio Code"]),
    ("Windows Terminal", "WindowsTerminal.exe", "Coding", ["pytest -q", "PowerShell", "ollama run qwen3:27b"]),
    ("Firefox", "firefox.exe", "Browsing", ["DuckDB – read_csv options — Mozilla Firefox", "LinkedIn Empleos — Mozilla Firefox",
                                             "YouTube — Tutorial de costura — Mozilla Firefox", "GitHub — Mozilla Firefox"]),
    ("Obsidian", "Obsidian.exe", "Writing", ["Capítulo 14 - Novela (borrador) - Obsidian", "Notas de mundo - Obsidian"]),
    ("Discord", "Discord.exe", "Communication", ["#fan-films - Discord", "#local-llm - Discord"]),
    ("Teams", "ms-teams.exe", "Meetings", ["Entrevista técnica | Microsoft Teams", "Reunión semanal | Microsoft Teams"]),
    ("Steam", "steam.exe", "Games", ["Steam"]),
    ("VLC", "vlc.exe", "Media", ["fan_edit_final.mkv - VLC media player"]),
]


def write_funes(out: Path) -> None:
    spans = []
    sid = 1
    day = date(2026, 9, 7)
    for _ in range(14):
        t = datetime.combine(day, datetime.min.time()) + timedelta(hours=9, minutes=rng.randint(0, 40))
        end_of_day = datetime.combine(day, datetime.min.time()) + timedelta(hours=rng.choice([19, 21, 23]))
        weekend = day.weekday() >= 5
        while t < end_of_day:
            weights = [6, 2, 5, 3, 2, 1, 4, 2] if weekend else [30, 8, 14, 6, 4, 3, 1, 1]
            app, exe, cat, titles = rng.choices(APPS, weights=weights)[0]
            dur = rng.randint(20, 2400)
            spans.append({"id": sid, "start_ts": t.timestamp(), "end_ts": (t + timedelta(seconds=dur)).timestamp(),
                          "kind": "window", "app": app, "exe": exe, "title": rng.choice(titles), "pid": rng.randint(1000, 30000),
                          "category": cat, "project": "laplaces-hoard" if "laplaces" in titles[0] else None,
                          "rules_version": 3, "open": 0})
            sid += 1
            t += timedelta(seconds=dur + rng.randint(0, 300))
        day += timedelta(days=1)
    payload = {"spans": spans, "file_events": [], "commits": [], "exported_at": datetime(2026, 9, 21, 22, 0).timestamp()}
    (out / "funes-export.json").write_text(json.dumps(payload), encoding="utf-8")


# ---------------------------------------------------------------- Daguerre --

def write_daguerre(out: Path) -> None:
    path = out / "daguerre-library.sqlite"
    path.unlink(missing_ok=True)
    con = sqlite3.connect(path)
    con.execute("""CREATE TABLE photos (id TEXT PRIMARY KEY, path TEXT, size INTEGER, taken_at TEXT, make TEXT, model TEXT,
                   lens TEXT, f_number REAL, exposure_time TEXT, iso INTEGER, focal_length REAL, gps_lat REAL, gps_lon REAL,
                   city TEXT, country TEXT, caption TEXT, thumb BLOB)""")
    cams = [("Apple", "iPhone 15 Pro", None, [6.8, 24.0]), ("SONY", "ILCE-7M3", "FE 24-70mm F2.8 GM", [24, 35, 50, 70]),
            ("FUJIFILM", "X-T30", "XF35mmF1.4 R", [35.0])]
    t = datetime(2024, 3, 1, 10)
    for i in range(640):
        make, model, lens, focals = rng.choices(cams, weights=[60, 30, 10])[0]
        t += timedelta(hours=rng.expovariate(1 / 30))
        iso = rng.choice([50, 64, 100, 200, 400, 800, 1600, 3200]) if make != "Apple" else rng.choice([32, 50, 80, 125, 400, 1000])
        city = rng.choice(["Madrid", "Madrid", "Madrid", "Segovia", "Toledo", "Lisboa", None])
        con.execute("INSERT INTO photos VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
            f"p{i:05d}", f"D:\\Fotos\\{t:%Y}\\{t:%m}\\IMG_{i:04d}.JPG", rng.randint(900_000, 9_000_000),
            t.isoformat(timespec="seconds"), make, model, lens, rng.choice([1.4, 1.8, 2.8, 4.0, 5.6, 8.0]),
            rng.choice(["1/60", "1/125", "1/250", "1/1000"]), iso, rng.choice(focals), None, None, city,
            "España" if city not in ("Lisboa", None) else ("Portugal" if city else None),
            rng.choice([None, "atardecer en el templo de Debod", "rodaje de un cortometraje, escena 3", "acueducto"]),
            bytes(rng.getrandbits(8) for _ in range(64))))
    con.commit()
    con.close()


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO_ROOT / "data-uxtest" / "files"
    out.mkdir(parents=True, exist_ok=True)
    write_bank(out)
    write_jobs(out)
    write_writing(out)
    write_bench(out)
    write_funes(out)
    write_daguerre(out)
    for f in sorted(out.iterdir()):
        print(f"{f.stat().st_size:>10,}  {f.name}")


if __name__ == "__main__":
    main()
